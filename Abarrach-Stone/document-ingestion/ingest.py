#!/usr/bin/env python3
"""
Document Ingestion Pipeline
Parses documents from staging and extracts structured entities via Gemini.
"""

import os
import sqlite3
import json
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Optional
import hashlib

# Config
STAGING_DIR = "/data/openclaw/staging"
DB_PATH = "/data/openclaw/documents.db"
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# Category mapping from folder names
CATEGORY_MAP = {
    "irs": "finance",
    "medical": "medical",
    "contracts": "contract",
    "resumes": "employment",
    "jobs": "employment",
    "interviews": "employment",
    "interviewprep": "employment",
    "binance": "finance",
    "kraken": "finance",
    "crypto": "finance",
    "projects": "skills",
    "visual studio": "skills",
    "rma": "employment",
    "bills": "finance",
    "loandocs": "finance",
    "homel loans": "finance",
}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".rtf", ".txt", ".txt"}


def init_db():
    """Initialize SQLite database."""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            filepath TEXT NOT NULL UNIQUE,
            category TEXT,
            content TEXT,
            file_hash TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            processed_at TIMESTAMP
        )
    """)
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS entities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER REFERENCES documents(id) ON DELETE CASCADE,
            entity_type TEXT,
            entity_key TEXT,
            entity_value TEXT,
            confidence REAL DEFAULT 1.0
        )
    """)
    
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_category ON documents(category)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_entity_type ON entities(entity_type)")
    
    conn.commit()
    return conn


def get_file_hash(filepath: str) -> str:
    """Get MD5 hash of file."""
    hash_md5 = hashlib.md5()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()


def detect_category(filepath: str) -> Optional[str]:
    """Detect category from folder path."""
    path_parts = Path(filepath).parts
    for part in path_parts:
        part_lower = part.lower()
        for key, cat in CATEGORY_MAP.items():
            if key in part_lower:
                return cat
    return None


def extract_text(filepath: str) -> Optional[str]:
    """Extract text from various file formats."""
    ext = Path(filepath).suffix.lower()
    
    try:
        if ext == ".pdf":
            # Use pdftotext if available
            result = subprocess.run(
                ["pdftotext", "-layout", filepath, "-"],
                capture_output=True, text=True, timeout=60
            )
            return result.stdout if result.returncode == 0 else None
            
        elif ext == ".docx":
            # Use python-docx
            try:
                import docx
                doc = docx.Document(filepath)
                return "\n".join([p.text for p in doc.paragraphs])
            except ImportError:
                return None
                
        elif ext == ".xlsx":
            # Use openpyxl
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, data_only=True)
                text = []
                for sheet in wb:
                    for row in sheet.iter_rows():
                        text.extend([str(cell.value) for cell in row if cell.value])
                return "\n".join(text)
            except ImportError:
                return None
                
        elif ext == ".rtf":
            # Strip RTF formatting
            with open(filepath, "r", errors="ignore") as f:
                content = f.read()
            # Simple RTF strip
            import re
            content = re.sub(r'\\[a-z]+\d*\s?', '', content)
            content = re.sub(r'[{}]', '', content)
            return content
            
        elif ext == ".txt":
            with open(filepath, "r", errors="ignore") as f:
                return f.read()
                
    except Exception as e:
        print(f"Error extracting {filepath}: {e}")
    
    return None


def extract_entities_gemini(content: str, category: str) -> list:
    """Extract structured entities using Gemini API."""
    if not GEMINI_API_KEY:
        print("GEMINI_API_KEY not set")
        return []
    
    import requests
    
    prompt = f"""Extract structured information from this document (category: {category}).

Return a JSON array of entities in this exact format:
{{"type": "identity|skills|employment|finance|medical|contract", "key": "field_name", "value": "extracted_value", "confidence": 0.0-1.0}}

Only extract fields you are confident about. If no useful information, return [].

Document content:
{content[:8000]}"""
    
    try:
        response = requests.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}",
            json={"contents": [{"parts": [{"text": prompt}]}]},
            timeout=30
        )
        
        if response.status_code == 200:
            text = response.json()["candidates"][0]["content"]["parts"][0]["text"]
            # Extract JSON from response
            import re
            match = re.search(r'\[.*\]', text, re.DOTALL)
            if match:
                return json.loads(match.group())
    except Exception as e:
        print(f"Gemini extraction error: {e}")
    
    return []


def scan_staging():
    """Scan staging directory for documents."""
    documents = []
    
    for root, dirs, files in os.walk(STAGING_DIR):
        for file in files:
            ext = Path(file).suffix.lower()
            if ext in SUPPORTED_EXTENSIONS:
                filepath = os.path.join(root, file)
                documents.append({
                    "filename": file,
                    "filepath": filepath,
                    "category": detect_category(filepath)
                })
    
    return documents


def process_documents(conn):
    """Process all documents in staging."""
    cursor = conn.cursor()
    
    # Get all documents to process
    documents = scan_staging()
    print(f"Found {len(documents)} documents to process")
    
    processed = 0
    for doc in documents:
        # Check if already processed (by file hash)
        try:
            file_hash = get_file_hash(doc["filepath"])
        except:
            continue
        
        cursor.execute("SELECT id FROM documents WHERE file_hash = ?", (file_hash,))
        if cursor.fetchone():
            continue
        
        # Extract text
        content = extract_text(doc["filepath"])
        if not content:
            continue
        
        # Store document
        cursor.execute("""
            INSERT INTO documents (filename, filepath, category, content, file_hash, processed_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (doc["filename"], doc["filepath"], doc["category"], content[:50000], file_hash, datetime.now().isoformat()))
        
        doc_id = cursor.lastrowid
        
        # Extract entities
        if doc["category"]:
            entities = extract_entities_gemini(content, doc["category"])
            for entity in entities:
                cursor.execute("""
                    INSERT INTO entities (document_id, entity_type, entity_key, entity_value, confidence)
                    VALUES (?, ?, ?, ?, ?)
                """, (doc_id, entity.get("type"), entity.get("key"), entity.get("value"), entity.get("confidence", 1.0)))
        
        conn.commit()
        processed += 1
        print(f"Processed: {doc['filename']} ({processed}/{len(documents)})")
    
    return processed


if __name__ == "__main__":
    print("Initializing database...")
    conn = init_db()
    
    print("Processing documents...")
    count = process_documents(conn)
    
    print(f"Done! Processed {count} documents.")
    conn.close()
